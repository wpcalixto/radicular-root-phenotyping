#!/usr/bin/env python3
"""
Title:
    Multimodal integration
File type:
    Script
Purpose:
    Joins image descriptors, electrical variables, management, chronology and final outcomes (air-
    dry shoot and root mass from Massa_full.xlsx and their ratio) by rhizotron and collection.
Repository:
    Software_Radicular_v1.0.0
Version:
    v1.0.0 (distribution); historical source dated 2026-07-13
Date:
    2026-07-13
Developer:
    Prof. Wesley Pacheco Calixto
Authors:
    Wesley Pacheco Calixto; Universidade de Coimbra (UC), Portugal; Universidade Federal de Goiás
    (UFG), Brazil; Instituto Federal de Educação, Ciência e Tecnologia de Goiás (IFG), Brazil
Corresponding author:
    Wesley Pacheco Calixto
License:
    MIT (see LICENSE)
Citation:
    If you use this code, please cite the Radicular Software record on Zenodo (DOI 10.5281/zenodo.23679263) and the related publication.
Related publication:
    Auditable longitudinal root phenotyping in rhizotrons using Frangi--U-Net fusion and apparent
    conductivity (manuscript in preparation for submission to Royal Society Open Science)
Inputs:
    - metricas/electrical_features_by_rhizotron.csv
    - metricas/root_features_by_image.csv
    - metricas/root_harvested.csv
    - tabdatas/Massa_full.xlsx
Outputs:
    - metricas/multimodal_long.csv
    - metricas/multimodal_targets.csv
    - metricas/multimodal_wide.csv
Dependencies:
    - openpyxl
    - local: config
Execution:
    python scripts/11_integracao.py (from a workspace prepared with tools/prepare_workspace.py)
Portability notes:
    Paths are resolved relative to the repository root through scripts/config.py; build the expected
    tree with tools/prepare_workspace.py.
Reproducibility notes:
    Verification in v1.0.0: re-executed in an isolated workspace built from the Dataset; outputs
    byte-identical to the archived versions. Environment of the verification: Python 3.9.6 (macOS)
    with the versions listed in requirements.txt.
Warnings:
    None specific to this file.
"""
# Original module notes (Portuguese) are preserved below.
"""
Fase 11 — Integração multimodal do dataset.

Une, por rizotron × coleta: descritores de imagem (g, Fase 09), variáveis elétricas
(e, Fase 10), variáveis de manejo (m: tratamento, adubação), cronologia (t: DAP) e os
alvos biológicos finais (y: massa aérea M_a, massa radicular M_r, razão R_a/r, e
comprimento radicular final L_r — proxy = L_skel da última coleta imageada).

Biomassa (`Massa_full.xlsx`): colunas {TR,T1,T2,T3}MA (massa aérea) e {TR,T1,T2,T3}MR
(massa radicular), 3 réplicas → 12 rizotrons (réplica r, tratamento t → R{t*3+r}).

Saídas:
  metricas/multimodal_long.csv   (1 linha por rizotron×coleta; g+e+m+t+y)
  metricas/multimodal_wide.csv   (1 linha por rizotron; features por coleta + y)
  metricas/multimodal_targets.csv (1 linha por rizotron; só os alvos y)
"""
import os
import sys
import csv
from collections import defaultdict

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import config as C

ROOTCSV = C.METRICAS / "root_features_by_image.csv"
ELECSV = C.METRICAS / "electrical_features_by_rhizotron.csv"
MASSA = C.TABDATAS / "Massa_full.xlsx"
TRATS = ["TR", "T1", "T2", "T3"]                  # prefixos na planilha (TR = REF)
COD = {"TR": "REF", "T1": "T1", "T2": "T2", "T3": "T3"}

# descritores de imagem e elétricos usados na integração
G_COLS = ["A_px", "O_ocupacao", "Dmax_norm", "Wmax_norm", "L_skel_px", "B_ramif", "N_br", "F_hull"]
E_COLS = ["R_ohm", "rho_a", "sigma_a"]


def fnum(x):
    """
    Purpose:
        Convert a CSV field to float, returning None for empty or non-numeric values.
    Parameters:
        x: value read from a CSV field
    Returns:
        float or None
    Raises:
        No exception is raised explicitly; errors of called libraries propagate.
    Notes:
        Parameter types are not annotated in the original code.
    """
    try:
        return float(x)
    except (TypeError, ValueError):
        return None


def ler_biomassa():
    """Retorna {rizotron: {'M_a':, 'M_r':, 'R_ar':}}."""
    ws = openpyxl.load_workbook(MASSA, data_only=True).active
    hdr = {str(ws.cell(1, c).value).strip().upper(): c for c in range(1, ws.max_column + 1)}
    out = {}
    for ti, t in enumerate(TRATS):
        for rep in (1, 2, 3):
            riz = f"R{ti * 3 + rep:02d}"
            Ma = ws.cell(1 + rep, hdr[f"{t}MA"]).value
            Mr = ws.cell(1 + rep, hdr[f"{t}MR"]).value
            Ma, Mr = fnum(Ma), fnum(Mr)
            out[riz] = {"M_a": Ma, "M_r": Mr,
                        "R_ar": round(Ma / Mr, 4) if (Ma and Mr) else ""}
    return out


def main():
    """
    Purpose:
        Run the script end to end: read inputs, compute results and write the outputs listed in the header.
    Parameters:
        none (some scripts accept an optional limit or command-line argument)
    Returns:
        None
    Raises:
        No exception is raised explicitly; errors of called libraries propagate.
    Notes:
        Parameter types are not annotated in the original code.
    """
    root = list(csv.DictReader(open(ROOTCSV, encoding="utf-8")))
    elec = {(r["rhizotron_id"], r["collection_date"]): r
            for r in csv.DictReader(open(ELECSV, encoding="utf-8"))}
    bio = ler_biomassa()

    # L_r final (proxy): L_skel na última coleta imageada de cada rizotron
    ult = {}
    for r in root:
        riz, dap = r["rizotron"], int(r["DAP"]) if r["DAP"] else -1
        if riz not in ult or dap > ult[riz][0]:
            ult[riz] = (dap, fnum(r["L_skel_px"]))
    L_r = {riz: v[1] for riz, v in ult.items()}

    # ---------- LONGO ----------
    long_rows = []
    for r in root:
        riz, data = r["rizotron"], r["data"]
        cod, nome = C.tratamento(riz)
        adub = C.ADUBACAO_PREVISTA.get(cod)          # DAP previsto de adubação (None p/ REF)
        dap = int(r["DAP"]) if r["DAP"] else ""
        e = elec.get((riz, data), {})
        b = bio.get(riz, {})
        linha = {
            "rizotron": riz, "tratamento": cod, "replicata": C.replicata(riz),
            "data": data, "DAP": dap, "coleta": r["coleta"],
            "adubacao_dap": adub if adub is not None else "",
            "dias_desde_adubacao": (dap - adub) if (adub is not None and dap != "") else "",
        }
        for c in G_COLS:
            linha[c] = r.get(c, "")
        for c in E_COLS:
            linha[c] = e.get(c, "")                  # vazio na coleta 7 (sem elétrica)
        linha["tem_eletrica"] = 1 if e else 0
        linha.update({"y_M_a": b.get("M_a", ""), "y_M_r": b.get("M_r", ""),
                      "y_R_ar": b.get("R_ar", ""), "y_L_r_final": L_r.get(riz, "")})
        long_rows.append(linha)
    long_rows.sort(key=lambda x: (x["rizotron"], x["coleta"] if x["coleta"] != "" else 0))

    campos_long = (["rizotron", "tratamento", "replicata", "data", "DAP", "coleta",
                    "adubacao_dap", "dias_desde_adubacao"] + G_COLS + E_COLS +
                   ["tem_eletrica", "y_M_a", "y_M_r", "y_R_ar", "y_L_r_final"])
    with open(C.METRICAS / "multimodal_long.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=campos_long); w.writeheader()
        for r in long_rows:
            w.writerow({k: r.get(k, "") for k in campos_long})

    # ---------- LARGO (1 linha por rizotron; features por coleta) ----------
    por_riz = defaultdict(dict)
    for r in long_rows:
        por_riz[r["rizotron"]][r["coleta"]] = r
    # colunas de feature por coleta (c1..c7) — descritores-chave de imagem + elétrica
    KEY_G = ["L_skel_px", "O_ocupacao", "A_px", "N_br"]
    KEY_E = ["sigma_a", "rho_a"]
    coletas = sorted({r["coleta"] for r in long_rows if r["coleta"] != ""}, key=lambda x: int(x))
    wide_rows = []
    for riz in sorted(por_riz):
        cod, _ = C.tratamento(riz)
        linha = {"rizotron": riz, "tratamento": cod, "replicata": C.replicata(riz)}
        for ck in coletas:
            src = por_riz[riz].get(ck, {})
            for g in KEY_G:
                linha[f"{g}_c{ck}"] = src.get(g, "")
            for e in KEY_E:
                linha[f"{e}_c{ck}"] = src.get(e, "")
        b = bio.get(riz, {})
        linha.update({"y_M_a": b.get("M_a", ""), "y_M_r": b.get("M_r", ""),
                      "y_R_ar": b.get("R_ar", ""), "y_L_r_final": L_r.get(riz, "")})
        wide_rows.append(linha)
    with open(C.METRICAS / "multimodal_wide.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(wide_rows[0].keys())); w.writeheader(); w.writerows(wide_rows)

    # raiz colhida (Fase 09B) — referência complementar de comprimento final.
    # Aqui o cm É legítimo: a raiz colhida foi fotografada COM trena (régua no quadro).
    harv = {}
    hp = C.METRICAS / "root_harvested.csv"
    if hp.exists():
        for r in csv.DictReader(open(hp, encoding="utf-8")):
            harv[r["rizotron"]] = (r.get("skel_len_px", ""), r.get("L_r_skel_cm", ""))

    # ---------- alvos ----------
    # L_r_final_insitu_px: proxy in-situ (L_skel da última coleta imageada) — SEM escala (px).
    # L_r_harvested_cm: referência com trena (régua real no quadro) — cm aqui é auditável.
    with open(C.METRICAS / "multimodal_targets.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["rizotron", "tratamento", "M_a", "M_r", "R_ar", "L_r_final_insitu_px",
                    "L_r_harvested_px", "L_r_harvested_cm"])
        for riz in sorted(bio):
            b = bio[riz]; cod, _ = C.tratamento(riz)
            hpx, hcm = harv.get(riz, ("", ""))
            w.writerow([riz, cod, b["M_a"], b["M_r"], b["R_ar"], L_r.get(riz, ""), hpx, hcm])

    import statistics as st
    print(f"[Fase 11] LONGO: {len(long_rows)} linhas (rizotron×coleta) -> multimodal_long.csv")
    print(f"   LARGO: {len(wide_rows)} rizotrons -> multimodal_wide.csv")
    print(f"   coletas com elétrica: {sum(1 for r in long_rows if r['tem_eletrica'])}/{len(long_rows)}")
    Rar = [b["R_ar"] for b in bio.values() if isinstance(b["R_ar"], (int, float))]
    print(f"   alvos: M_a e M_r de 12 rizotrons | R_a/r médio {st.mean(Rar):.2f} (min {min(Rar):.1f} max {max(Rar):.1f})")


if __name__ == "__main__":
    main()
